#!/usr/bin/env python3
"""Build PPTX presentation and take-home toolkit PDF for Cedric Milles."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt
from reportlab.lib.colors import HexColor, white, black
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm, mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
    KeepTogether,
    HRFlowable,
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "materiales"
OUT.mkdir(exist_ok=True)

NAVY = RGBColor(0x09, 0x2D, 0x4A)
BLUE = RGBColor(0x31, 0x5F, 0x8D)
SKY = RGBColor(0x4A, 0x8B, 0xC7)
GRAY = RGBColor(0x53, 0x61, 0x6F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xE8, 0xF1, 0xF8)
WARM = RGBColor(0xC4, 0x7A, 0x3A)

SLIDES = [
    {
        "kind": "dark",
        "block": "Portada",
        "title": "¿Quién decide cuánto valgo?",
        "body": "Autoestima, comparación, presión social y bullying\nFamilias de niños y niñas de 6 a 11 años\nComprender. Decidir. Avanzar.",
        "notes": "Dar la bienvenida. Presentar marca y duración: 40 min explicación + 20 min taller.",
    },
    {
        "kind": "content",
        "block": "0–4 min · Apertura",
        "title": "Cómo vamos a trabajar hoy",
        "bullets": [
            "Comprender: cómo se construye la autoestima entre 6 y 11 años",
            "Decidir: qué respuesta adulta ayuda ante error, comparación o burlas",
            "Avanzar: dos talleres y una herramienta para casa",
        ],
        "notes": "Explicar que no es terapia ni evaluación de familias.",
    },
    {
        "kind": "quote",
        "block": "Participación",
        "title": "Una pregunta para empezar",
        "quote": "Cuando vuestro hijo o hija dice «no valgo» o «soy tonto/a», ¿qué suele salir primero de vuestra boca?",
        "notes": "30–45 s en silencio. No pedir compartir.",
    },
    {
        "kind": "dark",
        "block": "Mensaje conductor",
        "title": "El valor no es una nota",
        "body": "La dignidad de un niño no depende de una calificación, un error, una comparación, su aspecto o la aceptación del grupo.\n\nSí podemos ayudar a que su manera de valorarse crezca con apoyo y vínculos seguros.",
        "notes": "Tesis central. Evitar promesas de cambio garantizado.",
    },
    {
        "kind": "content",
        "block": "Objetivos",
        "title": "Qué os llevaréis",
        "bullets": [
            "Distinguir autoestima, autoconcepto y autoeficacia",
            "Reformular una etiqueta global en una conducta concreta",
            "Acompañar con emoción + límite + paso posible",
            "Diferenciar crítica, conflicto y sospecha de acoso",
        ],
        "notes": "Objetivos observables, no clínicos.",
    },
    {
        "kind": "content",
        "block": "4–11 min · Conceptos",
        "title": "Tres ideas que no son lo mismo",
        "bullets": [
            "Autoestima: cómo me trato y cuánto valor me reconozco cuando algo sale mal",
            "Autoconcepto: cómo me describo en distintas áreas",
            "Autoeficacia: creer que puedo manejar esta tarea concreta",
        ],
        "notes": "Usar ejemplo de divisiones para autoeficacia.",
    },
    {
        "kind": "two",
        "block": "Conceptos",
        "title": "Una autoestima sana…",
        "left_title": "No es",
        "left": ["Sentirse superior", "No equivocarse nunca", "Aprobación constante", "Frases mágicas"],
        "right_title": "Sí es",
        "right": ["Realista y flexible", "Compatible con límites", "Capaz de pedir ayuda", "Separa persona y conducta"],
        "notes": "Autoestima frágil/contingente en una frase.",
    },
    {
        "kind": "two",
        "block": "Conceptos",
        "title": "Cuando el valor parece condicionado",
        "left_title": "Riesgo",
        "left": ["«Te quiero… cuando lo haces bien»", "Afecto usado como moneda"],
        "right_title": "Protección",
        "right": ["«Te quiero. Esta conducta hay que repararla»", "Pertenencia estable + responsabilidad"],
        "notes": "Patrones repetidos, no una frase aislada.",
    },
    {
        "kind": "quote",
        "block": "Micro · 45 s",
        "title": "«No sé hacer divisiones»",
        "quote": "¿Es baja autoestima, autoconcepto o autoeficacia?\nSuele ser autoeficacia: práctica y pasos, sin convertirla en «no valgo».",
        "notes": "Pedir mano alzada o pensar en silencio.",
    },
    {
        "kind": "two",
        "block": "11–16 min · Desarrollo",
        "title": "No es la misma etapa",
        "left_title": "6–8 años",
        "left": ["Frases breves y ejemplos", "Más co-regulación", "Semáforo y dibujos ayudan"],
        "right_title": "9–11 años",
        "right": ["Más peso del grupo", "Pueden ocultar malestar", "Preguntar qué ayuda necesitan"],
        "notes": "Edades orientativas.",
    },
    {
        "kind": "quote",
        "block": "Desarrollo",
        "title": "Primero calma, después conversación",
        "quote": "La regulación empieza muchas veces como co-regulación: prestamos calma, lenguaje y estructura.",
        "notes": "Conectar → bajar intensidad → orientar.",
    },
    {
        "kind": "content",
        "block": "16–24 min · Mecanismos",
        "title": "El ciclo que conviene ralentizar",
        "bullets": [
            "Situación: no me eligen para jugar",
            "Pensamiento: «nadie me quiere»",
            "Emoción: vergüenza, rabia",
            "Acción: aislarse o insultar",
            "Preguntas: ¿qué ocurrió? ¿qué me digo? ¿qué siento? ¿qué paso seguro?",
        ],
        "notes": "Diagrama oral si no hay animación.",
    },
    {
        "kind": "two",
        "block": "Mecanismos",
        "title": "Las etiquetas convierten conducta en identidad",
        "left_title": "«Eres vago»",
        "left": ["Identidad global", "Vergüenza", "Menos margen de cambio"],
        "right_title": "«Te cuesta empezar»",
        "right": ["Conducta concreta", "Sistema", "Primer paso posible"],
        "notes": "También etiquetas positivas pueden encerrar.",
    },
    {
        "kind": "content",
        "block": "Mecanismos",
        "title": "El feedback útil informa",
        "bullets": [
            "Observa: qué se hizo de forma concreta",
            "Significa: qué efecto tuvo",
            "Abre: cuál puede ser el siguiente paso",
            "Mejor «Has revisado el último paso» que «Eres un genio»",
        ],
        "notes": "Elogio de proceso honesto, no vacío.",
    },
    {
        "kind": "content",
        "block": "Mecanismos",
        "title": "El error informa si el reto es manejable",
        "bullets": [
            "Demasiado difícil → indefensión",
            "Todo resuelto por el adulto → poca eficacia",
            "Zona útil: desafío moderado + apoyo temporal + volver a probar",
        ],
        "notes": "Andamiaje / zona de desarrollo próximo en lenguaje cotidiano.",
    },
    {
        "kind": "two",
        "block": "Micro · 60 s",
        "title": "«He suspendido»",
        "left_title": "A",
        "left": ["«No pasa nada»"],
        "right_title": "B",
        "right": ["«La nota da información, no mide tu valor. Cuando baje el enfado, miramos qué falta.»"],
        "notes": "Pedir preferencia con gesto.",
    },
    {
        "kind": "content",
        "block": "24–29 min · Autonomía",
        "title": "Tres necesidades que orientan",
        "bullets": [
            "Autonomía: elecciones reales dentro de límites claros",
            "Competencia: sentir que puedo aprender y producir efectos",
            "Vínculo: sentirse cuidado sin condicionarlo al rendimiento",
        ],
        "notes": "Autodeterminación en versión familiar.",
    },
    {
        "kind": "quote",
        "block": "Autonomía",
        "title": "Validar no es ceder el límite",
        "quote": "«Tiene sentido que estés enfadado. No voy a dejar que rompas el material. Paramos dos minutos o hacemos juntos el primer paso.»",
        "notes": "Emoción + norma + alternativa.",
    },
    {
        "kind": "content",
        "block": "Autonomía",
        "title": "Frases listas para el momento",
        "bullets": [
            "«Ahora mismo te estás definiendo por esta tarea.»",
            "«¿Quieres que te escuche, pensemos opciones o intervengamos?»",
            "«No tienes que poder a solas; podemos dividirlo.»",
            "«Esto necesita reparación; sigo contigo.»",
        ],
        "notes": "Son propuestas, no frases mágicas.",
    },
    {
        "kind": "content",
        "block": "29–37 min · Comparación y acoso",
        "title": "«Todos son mejores»",
        "bullets": [
            "Suele comunicar evaluación, miedo a quedar fuera y deseo de reconocimiento",
            "Preguntar: ¿en qué te comparas? ¿con quién? ¿qué conclusión sacas?",
            "Ampliar criterios: progreso, estrategia, cuidado, cooperación",
        ],
        "notes": "No ridiculizar la importancia del grupo.",
    },
    {
        "kind": "content",
        "block": "Comparación y acoso",
        "title": "No todo conflicto es acoso",
        "bullets": [
            "Crítica útil: dato + impacto + posibilidad de mejora",
            "Conflicto: desacuerdo entre personas con poder similar",
            "Acoso: agresión no deseada, repetición y desequilibrio de poder",
        ],
        "notes": "No minimizar daños puntuales graves.",
    },
    {
        "kind": "content",
        "block": "Comparación y acoso",
        "title": "Ante sospecha de acoso",
        "bullets": [
            "Creer el malestar y escuchar",
            "Registrar hechos; conservar pruebas digitales",
            "Proteger: no dejar la solución solo al niño",
            "Coordinar con el protocolo del centro",
        ],
        "notes": "Responsabilidad adulta y sistémica.",
    },
    {
        "kind": "two",
        "block": "Micro · 45 s",
        "title": "«Me llaman tonto todos los recreos»",
        "left_title": "Evitar",
        "left": ["«No les hagas caso»", "«Sé más fuerte»"],
        "right_title": "Proteger",
        "right": ["«Gracias por contármelo»", "Apuntar hechos y pedir ayuda al colegio"],
        "notes": "Cerrar el bloque 5.",
    },
    {
        "kind": "content",
        "block": "37–40 min · Herramientas",
        "title": "La brújula del valor en casa",
        "bullets": [
            "Centro: «Mi valor no depende de este momento»",
            "Lo que siento · Lo que pasó · Lo que puedo hacer · A quién pido apoyo",
            "Ejemplo: no me han elegido para jugar",
        ],
        "notes": "Demostración rápida oral.",
    },
    {
        "kind": "content",
        "block": "Herramientas",
        "title": "Semáforo de calma",
        "bullets": [
            "Rojo: parar lo que puede hacer daño y bajar activación",
            "Ámbar: nombrar emoción, hechos e interpretación; elegir ayuda",
            "Verde: acción pequeña, hacerla y revisar el efecto",
        ],
        "notes": "Semáforo regula; brújula orienta.",
    },
    {
        "kind": "dark",
        "block": "Taller · 20 min",
        "title": "Pasamos a practicar",
        "body": "Trabajamos con casos ficticios.\nNo se evalúa a ninguna familia.\nSe puede pasar o trabajar en solitario.\n\nTaller 1 · Laboratorio de frases\nTaller 2 · Ensayo de acompañamiento",
        "notes": "Minutos 40–42: consignas y parejas.",
    },
    {
        "kind": "content",
        "block": "Taller 1 · 42–48",
        "title": "Laboratorio de frases",
        "bullets": [
            "Elegid 2 tarjetas y reescribidlas",
            "Incluir: validación + lo concreto + un paso o apoyo",
            "Tarjetas: No pasa nada / Eres el mejor / Tu hermano sí puede / Si quisieras podrías / No seas tan sensible / No les hagas caso",
        ],
        "notes": "Devolución breve de 1–2 ejemplos.",
    },
    {
        "kind": "two",
        "block": "Ejemplo",
        "title": "De etiqueta a acompañamiento",
        "left_title": "«No pasa nada»",
        "left": ["Minimiza el dolor"],
        "right_title": "Versión útil",
        "right": ["«Sí importa y entiendo que duela. Cuéntame qué pasó y vemos un primer paso.»"],
        "notes": "Mostrar pauta de revisión.",
    },
    {
        "kind": "two",
        "block": "Taller 2 · 48–55",
        "title": "Ensayo de acompañamiento",
        "left_title": "Ruta A · Error",
        "left": ["Pausa", "Validar", "Explorar después", "Primer paso"],
        "right_title": "Ruta B · Burlas",
        "right": ["Escuchar", "Seguridad", "Agradecer", "Registrar y pedir ayuda al centro"],
        "notes": "Intercambiar papeles. No representar humillaciones.",
    },
    {
        "kind": "content",
        "block": "Casos",
        "title": "Textos para el ensayo",
        "bullets": [
            "Ruta A: «Hoy he suspendido mates. Soy tonto. No quiero estudiar más.»",
            "Ruta B: «En el recreo me llaman nombres casi todos los días. No quiero ir al cole.»",
            "Nadie cuenta historias personales sensibles en grupo",
        ],
        "notes": "Si alguien se abre: agradecer y ofrecer hablar después.",
    },
    {
        "kind": "content",
        "block": "55–58 · Plan",
        "title": "Vuestro plan de 7 días",
        "bullets": [
            "Una situación concreta",
            "Una frase automática a cambiar",
            "Una práctica breve de 5–10 minutos",
            "Una revisión sin puntuar el valor del niño",
        ],
        "notes": "Compromiso pequeño y viable.",
    },
    {
        "kind": "dark",
        "block": "58–60 · Cierre",
        "title": "Una voz interior más útil",
        "body": "No buscamos «soy el mejor».\nBuscamos: «esto me afecta, puedo entenderlo, sigo teniendo valor y no tengo que resolverlo a solas».\n\nCedric Milles · cedricmilles.com",
        "notes": "Evaluación final: frase elegida + práctica de la semana.",
    },
    {
        "kind": "content",
        "block": "Cierre",
        "title": "Dos decisiones pequeñas",
        "bullets": [
            "Si mañana oís «soy tonto/a», ¿cuál será vuestra primera frase?",
            "¿Qué herramienta probaréis: brújula, semáforo o laboratorio de frases?",
            "Gracias por cuidar el valor de vuestros hijos e hijas con precisión y cariño",
        ],
        "notes": "Cierre cálido y breve.",
    },
]


def _set_run(run, *, size=24, bold=False, color=NAVY, font="Manrope"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _add_bg(slide, color):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _accent_bar(slide, color=SKY):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), Inches(7.5)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def _add_text_box(slide, left, top, width, height, text, *, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color)
    return box


def _add_bullets(slide, left, top, width, height, bullets, *, size=20, color=NAVY):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(10)
        for run in p.runs:
            _set_run(run, size=size, color=color)
    return box


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    logo = ROOT / "assets" / "isotipo.png"

    for data in SLIDES:
        slide = prs.slides.add_slide(blank)
        kind = data["kind"]
        if kind == "dark":
            _add_bg(slide, NAVY)
            _accent_bar(slide, WARM)
            _add_text_box(slide, Inches(0.7), Inches(0.55), Inches(11.5), Inches(0.4), data["block"], size=14, bold=True, color=SKY)
            _add_text_box(slide, Inches(0.7), Inches(1.3), Inches(11.8), Inches(1.8), data["title"], size=40, bold=True, color=WHITE)
            body = data.get("body", "")
            _add_text_box(slide, Inches(0.7), Inches(3.3), Inches(11.5), Inches(3.2), body, size=22, color=RGBColor(0xD7, 0xE8, 0xF7))
            if logo.exists() and data is SLIDES[0]:
                slide.shapes.add_picture(str(logo), Inches(11.6), Inches(0.4), height=Inches(0.9))
        elif kind == "quote":
            _add_bg(slide, WHITE)
            _accent_bar(slide, SKY)
            mist = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(2.2), Inches(11.9), Inches(3.2))
            mist.fill.solid()
            mist.fill.fore_color.rgb = MIST
            mist.line.fill.background()
            _add_text_box(slide, Inches(0.7), Inches(0.45), Inches(11.5), Inches(0.4), data["block"], size=14, bold=True, color=BLUE)
            _add_text_box(slide, Inches(0.7), Inches(1.0), Inches(11.8), Inches(1.0), data["title"], size=34, bold=True, color=NAVY)
            _add_text_box(slide, Inches(1.1), Inches(2.5), Inches(11.1), Inches(2.6), data["quote"], size=24, color=NAVY)
        elif kind == "two":
            _add_bg(slide, WHITE)
            _accent_bar(slide, SKY)
            _add_text_box(slide, Inches(0.7), Inches(0.45), Inches(11.5), Inches(0.4), data["block"], size=14, bold=True, color=BLUE)
            _add_text_box(slide, Inches(0.7), Inches(1.0), Inches(11.8), Inches(1.0), data["title"], size=32, bold=True, color=NAVY)
            left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(2.3), Inches(5.6), Inches(4.2))
            left.fill.solid()
            left.fill.fore_color.rgb = RGBColor(0xF8, 0xEC, 0xE8)
            left.line.fill.background()
            right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(2.3), Inches(5.6), Inches(4.2))
            right.fill.solid()
            right.fill.fore_color.rgb = RGBColor(0xE7, 0xF3, 0xEE)
            right.line.fill.background()
            _add_text_box(slide, Inches(1.0), Inches(2.55), Inches(5.0), Inches(0.5), data["left_title"], size=20, bold=True, color=NAVY)
            _add_bullets(slide, Inches(1.0), Inches(3.2), Inches(5.0), Inches(3.0), data["left"], size=18)
            _add_text_box(slide, Inches(7.3), Inches(2.55), Inches(5.0), Inches(0.5), data["right_title"], size=20, bold=True, color=NAVY)
            _add_bullets(slide, Inches(7.3), Inches(3.2), Inches(5.0), Inches(3.0), data["right"], size=18)
        else:
            _add_bg(slide, WHITE)
            _accent_bar(slide, SKY)
            _add_text_box(slide, Inches(0.7), Inches(0.45), Inches(11.5), Inches(0.4), data["block"], size=14, bold=True, color=BLUE)
            _add_text_box(slide, Inches(0.7), Inches(1.0), Inches(11.8), Inches(1.1), data["title"], size=34, bold=True, color=NAVY)
            _add_bullets(slide, Inches(0.9), Inches(2.4), Inches(11.5), Inches(4.4), data.get("bullets", []), size=22)

        notes_tf = slide.notes_slide.notes_text_frame
        notes_tf.text = data.get("notes", "")

    path = OUT / "Quien_decide_cuanto_valgo_Presentacion.pptx"
    prs.save(path)
    return path


def build_toolkit_pdf():
    path = OUT / "Quien_decide_cuanto_valgo_Kit_familias.pdf"
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=1.6 * cm,
        rightMargin=1.6 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    navy = HexColor("#092D4A")
    blue = HexColor("#315F8D")
    sky = HexColor("#4A8BC7")
    gray = HexColor("#53616F")

    title = ParagraphStyle("T", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, textColor=navy, spaceAfter=8, leading=24)
    h = ParagraphStyle("H", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, textColor=blue, spaceBefore=12, spaceAfter=6)
    body = ParagraphStyle("B", parent=styles["Normal"], fontName="Helvetica", fontSize=10, textColor=navy, leading=14, spaceAfter=6)
    small = ParagraphStyle("S", parent=styles["Normal"], fontName="Helvetica", fontSize=9, textColor=gray, leading=12, spaceAfter=4)
    box = ParagraphStyle("X", parent=styles["Normal"], fontName="Helvetica", fontSize=10, textColor=navy, leading=13)

    story = []
    story.append(Paragraph("Cedric Milles · Psicología · Orientación · Rendimiento", small))
    story.append(Paragraph("¿Quién decide cuánto valgo?", title))
    story.append(Paragraph("Kit de herramientas para familias (6–11 años). Uso psicoeducativo; no sustituye valoración profesional.", body))
    story.append(HRFlowable(width="100%", thickness=1, color=sky, spaceAfter=10))

    story.append(Paragraph("1. Mi brújula para acompañar", h))
    story.append(Paragraph("<b>Centro:</b> «Mi valor no depende de este momento».", body))
    data = [
        [Paragraph("<b>Comprender</b><br/>Observar y escuchar. ¿Qué siente? ¿Qué hechos sé?", box),
         Paragraph("<b>Decidir</b><br/>¿Validar, poner límite o proteger?", box)],
        [Paragraph("<b>Avanzar</b><br/>Un paso o ayuda concreta + revisión breve.", box),
         Paragraph("<b>Ejemplo</b><br/>No le eligen → vergüenza → preguntar siguiente ronda o pedir apoyo adulto si se repite.", box)],
    ]
    t = Table(data, colWidths=[8.5 * cm, 8.5 * cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), HexColor("#E8F1F8")),
        ("BOX", (0, 0), (-1, -1), 0.5, sky),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, sky),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(t)

    story.append(Paragraph("2. Semáforo de calma", h))
    story.append(Paragraph("<b>Rojo:</b> parar lo dañino y bajar activación. <b>Ámbar:</b> nombrar emoción, hechos e interpretación; elegir ayuda. <b>Verde:</b> acción pequeña y revisar.", body))
    story.append(Paragraph("Señales corporales posibles: nudo en el estómago, cara caliente, puños, voz alta, ganas de huir. Edad 6–8: una pregunta por color. 9–11: distinguir dato / interpretación / valor.", small))

    story.append(Paragraph("3. Registro situación–pensamiento–emoción–acción", h))
    reg = [
        [Paragraph("<b>Situación</b>", box), Paragraph("________________________________", box)],
        [Paragraph("<b>Pensamiento</b>", box), Paragraph("________________________________", box)],
        [Paragraph("<b>Emoción / cuerpo</b>", box), Paragraph("________________________________", box)],
        [Paragraph("<b>Acción</b>", box), Paragraph("________________________________", box)],
        [Paragraph("<b>Alternativa útil</b>", box), Paragraph("________________________________", box)],
    ]
    rt = Table(reg, colWidths=[4.5 * cm, 12.5 * cm])
    rt.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.5, navy),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HexColor("#C9D7E5")),
        ("BACKGROUND", (0, 0), (0, -1), HexColor("#F4F8FB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(rt)
    story.append(Paragraph("Ejemplo: No me eligen → «nadie me quiere» → vergüenza → me voy → alternativa: preguntar o pedir ayuda si se repite. No usar en plena crisis.", small))

    story.append(PageBreak())
    story.append(Paragraph("4. Escudo de fortalezas (sin clasificar por notas)", h))
    esc = [
        [Paragraph("<b>Intereses / lo que me gusta</b><br/><br/><br/>", box),
         Paragraph("<b>Algo que he aprendido</b><br/><br/><br/>", box)],
        [Paragraph("<b>Un valor o conducta que practico</b><br/><br/><br/>", box),
         Paragraph("<b>Personas que me ayudan</b><br/><br/><br/>", box)],
    ]
    et = Table(esc, colWidths=[8.5 * cm, 8.5 * cm], rowHeights=[3.2 * cm, 3.2 * cm])
    et.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 1, navy),
        ("INNERGRID", (0, 0), (-1, -1), 0.6, sky),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(et)

    story.append(Paragraph("5. Reto gradual con apoyos", h))
    story.append(Paragraph("Objetivo pequeño: ____________________________", body))
    story.append(Paragraph("Paso 1 (con ayuda): __________  Paso 2 (juntos): __________  Paso 3 (solo con pista): __________", body))
    story.append(Paragraph("Revisión: ¿qué aprendí? (no solo si lo conseguí). Límite: no usar para exponer a violencia o acoso.", small))

    story.append(Paragraph("6. Filtro de opiniones", h))
    story.append(Paragraph("Cuando llega una crítica o comparación: 1) ¿Es un dato o una opinión? 2) ¿Quién la dice y con qué intención? 3) ¿Qué parte me sirve? 4) ¿Qué no mide sobre mi valor?", body))

    story.append(Paragraph("7. Tarjeta ante burlas o sospecha de acoso", h))
    story.append(Paragraph("1) Escuchar y creer el malestar. 2) Agradecer que lo cuente. 3) Registrar hechos. 4) Adultos de apoyo: familia / tutor / orientación. 5) Coordinar con el centro. No pedir al niño que lo detenga solo.", body))

    story.append(Paragraph("8. Frases recortables (adultos)", h))
    phrases = [
        "«Sí importa y entiendo que duela.»",
        "«¿Quieres que te escuche o pensemos un primer paso?»",
        "«Ahora mismo te estás definiendo por esta tarea.»",
        "«Esto necesita reparación; sigo contigo.»",
        "«La nota da información, no mide tu valor.»",
        "«Gracias por contármelo; vamos a pedir ayuda.»",
    ]
    for ph in phrases:
        story.append(Paragraph(f"□ {ph}", body))

    story.append(Paragraph("9. Plan de 7 días (una sola práctica)", h))
    plan = [
        [Paragraph("<b>Situación que quiero acompañar mejor</b>", box), Paragraph("_______________________________", box)],
        [Paragraph("<b>Frase automática que quiero cambiar</b>", box), Paragraph("_______________________________", box)],
        [Paragraph("<b>Práctica breve (5–10 min)</b>", box), Paragraph("_______________________________", box)],
        [Paragraph("<b>Momento de revisión (sin notas)</b>", box), Paragraph("_______________________________", box)],
    ]
    pt = Table(plan, colWidths=[7 * cm, 10 * cm])
    pt.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.5, navy),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HexColor("#C9D7E5")),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(pt)
    story.append(Spacer(1, 8))
    story.append(Paragraph("Cuándo pedir ayuda profesional: malestar intenso o persistente con impacto en sueño, colegio, relaciones o actividades; indicios de acoso o riesgo. En riesgo inmediato, priorizar seguridad y servicios de emergencia.", small))
    story.append(Paragraph("cedricmilles.com · Material psicoeducativo para la charla «¿Quién decide cuánto valgo?»", small))

    doc.build(story)
    return path


def build_one_pager():
    path = OUT / "Quien_decide_cuanto_valgo_Brujula_una_pagina.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.6 * cm, bottomMargin=1.6 * cm)
    styles = getSampleStyleSheet()
    navy = HexColor("#092D4A")
    blue = HexColor("#315F8D")
    title = ParagraphStyle("T", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, textColor=navy, spaceAfter=10)
    body = ParagraphStyle("B", parent=styles["Normal"], fontName="Helvetica", fontSize=11, textColor=navy, leading=15, spaceAfter=8)
    h = ParagraphStyle("H", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=blue, spaceBefore=8, spaceAfter=4)
    story = [
        Paragraph("Mi brújula para acompañar", title),
        Paragraph("Cedric Milles · Charla «¿Quién decide cuánto valgo?» · Familias 6–11 años", body),
        Paragraph("<b>Frase centro:</b> Mi valor no depende de este momento.", body),
        Paragraph("Comprender", h),
        Paragraph("Escucho sin interrumpir. Nombro la emoción con suavidad. Separar hechos y suposiciones.", body),
        Paragraph("Decidir", h),
        Paragraph("¿Hace falta validar, sostener un límite o proteger/coordinar con el colegio?", body),
        Paragraph("Avanzar", h),
        Paragraph("Acordamos un paso pequeño o una ayuda. Revisamos después sin puntuar el valor del niño o la niña.", body),
        Paragraph("Recordatorio", h),
        Paragraph("Validar ≠ ceder. Feedback específico. Ante acoso: proteger y activar adultos. Una práctica breve a la semana basta para empezar.", body),
    ]
    doc.build(story)
    return path


if __name__ == "__main__":
    pptx = build_pptx()
    kit = build_toolkit_pdf()
    one = build_one_pager()
    print(pptx)
    print(kit)
    print(one)
